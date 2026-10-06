#!/usr/bin/env python3
"""Offline regression/negative tests. The 17 KB fixture preserves official cells.

No test uses live HTTP or installs packages. Synthetic geometries are restricted
 to validation fixtures; full saved-source reproduction uses real BA Data bytes.
"""
import copy
import io
import json
import sys
import tempfile
import types
import unittest
from contextlib import ExitStack, redirect_stdout
from pathlib import Path
from unittest.mock import patch

# Import without requests installed, and prohibit accidental HTTP in all tests.
def no_network(*args, **kwargs):
    raise AssertionError('Network is forbidden in offline tests')
try:
    import requests
except ModuleNotFoundError:
    offline_requests = types.ModuleType('requests')
    offline_requests.get = no_network
    sys.modules['requests'] = offline_requests
from openpyxl import Workbook
import generar_estructura_actual_v2 as p
import verificar_estructura_actual as v

def setUpModule():
    global network_guard
    network_guard = patch.object(p.requests, 'get', side_effect=no_network)
    network_guard.start()


def tearDownModule():
    network_guard.stop()


FIXTURE = json.loads((Path(__file__).parent/'tests/fixtures/estructura_actual.json').read_text())
PDF_URL = FIXTURE['manifest'][2]['url']
PAGE = p.IDECBA_PUBLICATIONS + 'ejes-comerciales-ciudad-de-buenos-aires-2do-cuatrimestre-de-2026/'
HTML = '<h3>Ejes comerciales. Ciudad de Buenos Aires. 2do. cuatrimestre de 2026</h3><a href="' + PDF_URL + '">PDF</a>'


def book(key, q=2):
    wb = Workbook()
    wb.remove(wb.active)
    for item in FIXTURE[key]:
        if q == 1 and p.period(item['title']) == (2026, 2):
            continue
        ws = wb.create_sheet(item['title'])
        for row in item['rows']:
            ws.append(row)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7 if key == 'indicadores' else 17)
        if key == 'rubros':
            ws.merge_cells('A2:A3')
            ws.merge_cells('B2:B3')
            ws.merge_cells('C2:Q2')
    return wb


def as_bytes(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def geo_fixture():
    # Intentionally synthetic geometry, only for unit tests of the existing
    # identity/schema contract. Real geometries are used by run_offline_sources.
    return {'type': 'FeatureCollection', 'features': [
        {'type': 'Feature', 'properties': {'comuna': c, 'barrios': ''},
         'geometry': {'type': 'Polygon', 'coordinates': [[[0, 0], [1, 0], [0, 1], [0, 0]]]}}
        for c in range(1, 16)]}


def source_context(directory, iw=None, rw=None, cover=None, html=HTML, download_failure=False):
    stack = ExitStack()
    out, geo = Path(directory)/'actual.json', Path(directory)/'comunas.geojson'
    raw = {'https://www.estadisticaciudad.gob.ar/ind.xlsx': as_bytes(iw or book('indicadores')),
           'https://www.estadisticaciudad.gob.ar/rub.xlsx': as_bytes(rw or book('rubros')),
           PDF_URL: b'PDF reader isolated in this unit fixture', p.OEDE_URL: b'OEDE reader isolated',
           p.COMUNAS_URL: json.dumps(geo_fixture()).encode()}
    stack.enter_context(patch.object(p, 'OUT', out))
    stack.enter_context(patch.object(p, 'GEO_OUT', geo))
    stack.enter_context(patch.object(p, 'discover', return_value=([list(raw)[1]], [list(raw)[0]])))
    stack.enter_context(patch.object(p, 'get', side_effect=RuntimeError('download failed') if download_failure else raw.__getitem__))
    stack.enter_context(patch.object(p, 'get_text', return_value=html))
    page = types.SimpleNamespace(extract_text=lambda: FIXTURE['cover_c2'] if cover is None else cover)
    stack.enter_context(patch.object(p, 'PdfReader', return_value=types.SimpleNamespace(pages=[page])))
    stack.enter_context(patch.object(p, 'parse_oede', return_value=copy.deepcopy(FIXTURE['official_oede_output_c1'])))
    stack.enter_context(redirect_stdout(io.StringIO()))
    return stack


class ParserTests(unittest.TestCase):
    def parse(self, wb):
        return p.parse_indicadores(wb, p.parse_rubros(book('rubros'))[2], include_comparison=True)

    def test_semantic_header_ignores_merged_title(self):
        ws = p.latest_sheet(book('indicadores'))[1]
        self.assertIn('interanual', ws['A1'].value)
        self.assertEqual(p.semantic_column(ws, 'interanual'), 6)

    def test_c2_matches_all_official_cells(self):
        period, relevados, ocupados, tasa, comunas, comparison = self.parse(book('indicadores'))
        self.assertEqual((period, relevados, ocupados, tasa), ((2026, 2), 12916, 11528, 89.3))
        self.assertEqual(comparison['desde'], {'anio': 2025, 'cuatrimestre': 2})
        self.assertEqual(comparison['hasta'], {'anio': 2026, 'cuatrimestre': 2})
        self.assertEqual((comparison['tasa_ocupacion_desde'], comparison['variacion_total_pp']), (91.5, -2.2))
        current = p.latest_sheet(book('indicadores'))[1]
        prev = p.sheet_for(book('indicadores'), (2025, 2))
        for c in range(1, 16):
            self.assertEqual(comunas[str(c)]['variacion_interanual_pp'], round(current.cell(c + 3, 6).value, 1))
            self.assertEqual(comunas[str(c)]['tasa_ocupacion_anterior'], round(prev.cell(c + 3, 4).value, 1))

    def test_c2_rounding_previous_official_rates(self):
        comunas = self.parse(book('indicadores'))[4]
        self.assertEqual({c: comunas[c]['tasa_ocupacion_anterior'] for c in ('2', '4', '9', '12')}, {'2': 93.2, '4': 89.1, '9': 89.2, '12': 93.5})

    def test_c1_regression_only_justified_previous_rate_corrections(self):
        rw, iw = book('rubros', 1), book('indicadores', 1)
        _, _, occ, _ = p.parse_rubros(rw)
        result = p.parse_indicadores(iw, occ, include_comparison=True)
        expected = copy.deepcopy(FIXTURE['original_c1_comunas'])
        corrections = {'1': 87.4, '8': 88.8, '14': 89.9, '15': 89.9}
        for cid, value in corrections.items():
            self.assertNotEqual(expected[cid]['tasa_ocupacion_anterior'], value)
            expected[cid]['tasa_ocupacion_anterior'] = value
        self.assertEqual(result[:4], ((2026, 1), 12896, 11605, 90.0))
        self.assertEqual(result[4], expected)
        self.assertEqual((result[5]['tasa_ocupacion_desde'], result[5]['variacion_total_pp']), (91.6, -1.6))

    def test_missing_or_ambiguous_indicator_headers(self):
        for cell, value in [('F2', None), ('H2', 'Variación interanual (p.p.)'), ('B2', None), ('H2', 'Locales Relevados')]:
            with self.subTest(cell=cell, value=value):
                wb = book('indicadores')
                p.latest_sheet(wb)[1][cell] = value
                with self.assertRaises(RuntimeError): self.parse(wb)
        wb = book('indicadores'); ws = p.latest_sheet(wb)[1]
        for col in range(1, 8): ws.cell(10, col).value = ws.cell(2, col).value
        with self.assertRaisesRegex(RuntimeError, 'encabezado'): self.parse(wb)

    def test_all_missing_deltas_fail_not_silent_success(self):
        wb = book('indicadores'); ws = p.latest_sheet(wb)[1]
        for r in range(3, 19): ws.cell(r, 6).value = None
        with self.assertRaisesRegex(RuntimeError, 'interanual'): self.parse(wb)

    def test_invalid_indicators_fail(self):
        cases = [('F18', None), ('B18', None), ('A18', 14), ('A18', 16), ('A18', 1.4),
                 ('B4', 1910.4), ('C4', 1657), ('B3', 12917), ('D4', 86.9),
                 ('F3', -1.8), ('F4', -1.0), ('F4', float('nan')), ('F4', float('inf'))]
        for cell, value in cases:
            with self.subTest(cell=cell, value=value):
                wb = book('indicadores'); p.latest_sheet(wb)[1][cell] = value
                with self.assertRaises(RuntimeError): self.parse(wb)

    def test_duplicate_total_fails(self):
        wb = book('indicadores'); ws = p.latest_sheet(wb)[1]
        ws.append([ws.cell(3, c).value for c in range(1, 8)])
        with self.assertRaisesRegex(RuntimeError, 'duplicado'): self.parse(wb)

    def test_missing_previous_sheet_fails(self):
        wb = book('indicadores'); wb.remove(p.sheet_for(wb, (2025, 2)))
        with self.assertRaisesRegex(RuntimeError, 'ausente'): self.parse(wb)

    def test_ambiguous_sheet_fails(self):
        wb = book('indicadores'); copy_ws = wb.copy_worksheet(p.latest_sheet(wb)[1]); copy_ws.title = '2do. cuatrimestre de 2026 copia'
        with self.assertRaisesRegex(RuntimeError, 'ambigua'): self.parse(wb)

    def test_sheet_title_period_mismatch_fails(self):
        wb = book('indicadores'); p.latest_sheet(wb)[1]['A1'] = '48 ejes comerciales. 1er. cuatrimestre de 2026'
        with self.assertRaisesRegex(RuntimeError, 'período'): self.parse(wb)

    def test_previous_rates_must_reconcile(self):
        wb = book('indicadores'); p.sheet_for(wb, (2025, 2))['D4'] = 88.7
        with self.assertRaisesRegex(RuntimeError, 'tasa'): self.parse(wb)

    def test_rubros_reconcile_rows_and_columns(self):
        _, total, communes, rubros = p.parse_rubros(book('rubros'))
        self.assertEqual((total, len(rubros), len(communes)), (11528, 19, 15))
        self.assertTrue(all(sum(x['comunas'].values()) == x['total'] for x in rubros))
        self.assertTrue(all(sum(x['comunas'][c] for x in rubros) == communes[c] for c in communes))

    def test_rubros_missing_duplicated_or_inconsistent_fail(self):
        for cell, value in [('Q3', 14), ('Q3', None), ('Q5', None), ('Q5', -1), ('Q5', 234), ('B5', 3394), ('A6', 'Indumentaria, textiles y calzado')]:
            with self.subTest(cell=cell, value=value):
                wb = book('rubros'); p.latest_sheet(wb)[1][cell] = value
                with self.assertRaises(RuntimeError): p.parse_rubros(wb)
        wb = book('rubros'); p.latest_sheet(wb)[1].delete_rows(5)
        with self.assertRaises(RuntimeError): p.parse_rubros(wb)

    def test_rubros_column_totals_fail_even_if_rows_sum(self):
        wb = book('rubros'); ws = p.latest_sheet(wb)[1]
        ws['C5'] = ws['C5'].value + 1; ws['D5'] = ws['D5'].value - 1
        with self.assertRaisesRegex(RuntimeError, 'por comuna'): p.parse_rubros(wb)

    def test_crossbook_occupied_mismatch_fails(self):
        occ = p.parse_rubros(book('rubros'))[2]; occ['1'] += 1
        with self.assertRaisesRegex(RuntimeError, 'no coinciden'): p.parse_indicadores(book('indicadores'), occ)

    def test_period_does_not_read_index_sheet_as_period(self):
        self.assertIsNone(p.period('AC_EJ_2026_04'))
        self.assertIsNone(p.period('1er. cuatrimestre de 2025/2do. cuatrimestre de 2026'))
        self.assertEqual(p.period('2do-cuatrimestre-de-2026'), (2026, 2))


class PipelineTests(unittest.TestCase):
    def document(self):
        with tempfile.TemporaryDirectory() as tmp, source_context(tmp):
            p.main()
            return json.loads((Path(tmp)/'actual.json').read_text())

    def test_full_unit_generation_and_verification(self):
        d = self.document()
        self.assertIn('válido', v.validate(d, geo_fixture()))
        self.assertEqual(d['fuentes']['idecba_informe']['numero'], 2063)
        self.assertEqual(d['fuentes']['idecba_informe']['url'], PDF_URL)

    def test_report_wrong_period_title_or_pdf_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            for kwargs in [{'html': HTML.replace('2do.', '1er.')}, {'cover': FIXTURE['cover_c2'].replace('2do.', '1er.')},
                           {'cover': FIXTURE['cover_c2'].replace('2063', '2033')}, {'html': HTML.replace('href=', 'nohref=')},
                           {'html': HTML + '<a href="' + PDF_URL.replace('2063', '2033') + '">Other</a>'}]:
                with self.subTest(kwargs=kwargs), source_context(tmp, **kwargs):
                    with self.assertRaises(RuntimeError): p.discover_report((2026, 2))

    def test_report_duplicate_pdf_links_are_not_ambiguous(self):
        with tempfile.TemporaryDirectory() as tmp, source_context(tmp, html=HTML+HTML.split('</h3>')[1]):
            self.assertEqual(p.discover_report((2026, 2))['numero'], 2063)

    def test_failed_download_parser_period_report_or_verifier_preserves_last_bytes(self):
        invalid = book('indicadores'); p.latest_sheet(invalid)[1]['F18'] = None
        cases = [{'download_failure': True}, {'iw': invalid}, {'rw': book('rubros', 1)}, {'cover': 'wrong report'}]
        for kwargs in cases:
            with self.subTest(kwargs=list(kwargs)), tempfile.TemporaryDirectory() as tmp, source_context(tmp, **kwargs):
                out, geo = Path(tmp)/'actual.json', Path(tmp)/'comunas.geojson'
                out.write_bytes(b'LAST VALID JSON'); geo.write_bytes(b'LAST VALID GEO')
                with self.assertRaises(RuntimeError): p.main()
                self.assertEqual(out.read_bytes(), b'LAST VALID JSON')
                self.assertEqual(geo.read_bytes(), b'LAST VALID GEO')
        with tempfile.TemporaryDirectory() as tmp, source_context(tmp), patch.object(v, 'validate', side_effect=ValueError('invalid')):
            out = Path(tmp)/'actual.json'; out.write_bytes(b'LAST VALID')
            with self.assertRaises(ValueError): p.main()
            self.assertEqual(out.read_bytes(), b'LAST VALID')

    def test_old_period_cannot_overwrite_newer_valid_output(self):
        with tempfile.TemporaryDirectory() as tmp, source_context(tmp):
            out = Path(tmp)/'actual.json'
            raw = b'{"panorama":{"ejes_comerciales":{"periodo":{"anio":2026,"cuatrimestre":3}}}}'
            out.write_bytes(raw)
            with self.assertRaisesRegex(RuntimeError, 'retroceso'): p.main()
            self.assertEqual(out.read_bytes(), raw)

    def test_atomic_stage_and_replace_failure_preserves_both_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            one, two = Path(tmp)/'one.json', Path(tmp)/'two.json'
            one.write_bytes(b'ONE'); two.write_bytes(b'TWO')
            with self.assertRaises(ValueError): p.publish_json({one: {}, two: {'invalid': float('nan')}})
            self.assertEqual((one.read_bytes(), two.read_bytes()), (b'ONE', b'TWO'))
            original = p.os.replace
            calls = []
            def failure(src, dst):
                calls.append(str(dst))
                if len(calls) == 2: raise OSError('simulated disk failure')
                return original(src, dst)
            with patch.object(p.os, 'replace', side_effect=failure), self.assertRaises(OSError):
                p.publish_json({one: {}, two: {}})
            self.assertEqual((one.read_bytes(), two.read_bytes()), (b'ONE', b'TWO'))
            self.assertEqual(sorted(x.name for x in Path(tmp).iterdir()), ['one.json', 'two.json'])

    def test_verifier_rejects_missing_total_wrong_period_and_mislabelled_sources(self):
        original = self.document()
        mutations = [
            lambda d: d['panorama']['ejes_comerciales']['comparacion_interanual'].pop('variacion_total_pp'),
            lambda d: d['panorama']['ejes_comerciales']['comparacion_interanual'].__setitem__('tasa_ocupacion_desde', None),
            lambda d: d['panorama']['ejes_comerciales']['comparacion_interanual']['desde'].__setitem__('cuatrimestre', 1),
            lambda d: d['panorama']['ejes_comerciales']['comparacion_interanual']['hasta'].__setitem__('anio', 2025),
            lambda d: d['fuentes']['idecba_indicadores']['periodo'].__setitem__('cuatrimestre', 1),
            lambda d: d['fuentes']['idecba_informe'].__setitem__('url', PDF_URL.replace('2063', '2033')),
            lambda d: d['fuentes']['idecba_informe']['periodo_verificado'].__setitem__('cuatrimestre', 1),
            lambda d: d['fuentes']['idecba_informe'].__setitem__('pagina', PAGE.replace('2do', '1er')),
            lambda d: d['panorama']['ejes_comerciales']['comunas'].pop('15'),
            lambda d: d['panorama']['ejes_comerciales']['comunas']['1'].__setitem__('tasa_ocupacion', 87.0),
            lambda d: d['panorama']['ejes_comerciales']['rubros'][0]['comunas'].__setitem__('1', 254),
            lambda d: d['panorama']['ejes_comerciales']['rubros'].pop(),
            lambda d: d['panorama']['ejes_comerciales'].__setitem__('locales_relevados', 12917),
        ]
        for idx, mutate in enumerate(mutations):
            with self.subTest(case=idx):
                d = copy.deepcopy(original); mutate(d)
                with self.assertRaises(ValueError): v.validate(d, geo_fixture())

    def test_verifier_rejects_duplicate_json_keys(self):
        with self.assertRaisesRegex(ValueError, 'duplicada'):
            json.loads('{"comunas":{"1":{},"1":{}}}', object_pairs_hook=v.unique_object)


if __name__ == '__main__':
    unittest.main()
