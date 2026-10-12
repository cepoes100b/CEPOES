"""Regresiones EIM sobre los XLSX canónicos versionados, sin descargar fuentes."""
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import parsers as P
import verificar as V

ROOT = Path(__file__).resolve().parents[1]
FUENTES = ('industria_ing.xlsx', 'masa_salarial.xlsx')


class EncuestaIndustrialTest(unittest.TestCase):
    def libro(self, fuente):
        return openpyxl.load_workbook(ROOT / 'idecba' / fuente, data_only=True)

    def parsear_modificado(self, editar, fuente='industria_ing.xlsx'):
        wb = self.libro(fuente)
        editar(wb.worksheets[0])
        with patch.object(P, 'abrir', return_value=(wb, wb.worksheets[0])):
            return P._bloques_por_rama(fuente)

    def test_fuentes_reales_todos_los_valores_y_bloques(self):
        for fuente in FUENTES:
            with self.subTest(fuente=fuente):
                ws = self.libro(fuente).worksheets[0]
                periodos, bloques = P._bloques_por_rama(ROOT / 'idecba' / fuente)
                V.validar_periodos_mensuales(periodos)
                self.assertEqual(periodos[0], '2001-10')
                # El tramo auditado siempre se conserva aunque la fuente crezca.
                self.assertGreaterEqual(len(periodos), 297)
                self.assertEqual(periodos[296], '2026-06')
                for anio in range(2016, 2021):
                    fila_junio = 180 + (anio - 2016) * 12
                    self.assertEqual(P.anio_de(ws.cell(fila_junio, 1).value), anio)
                    for mes in range(1, 13):
                        r = fila_junio + mes - 6
                        self.assertEqual(periodos[r - 4], f'{anio}-{mes:02d}')
                inicios = [c for c in range(3, ws.max_column + 1)
                           if P.norm(ws.cell(3, c).value) == 'total']
                self.assertEqual(len(bloques), len(inicios))
                for i, (bloque, c) in enumerate(zip(bloques.values(), inicios)):
                    self.assertEqual(bloque['total'], [P.r1(P.num(ws.cell(r, c).value), 2)
                                                     for r in range(4, len(periodos) + 4)])
                    fin = inicios[i + 1] if i + 1 < len(inicios) else ws.max_column + 1
                    esperadas = {str(ws.cell(3, col).value).strip(): col
                                 for col in range(c + 1, fin) if ws.cell(3, col).value}
                    self.assertEqual(set(bloque['ramas']), set(esperadas))
                    for nombre, col in esperadas.items():
                        self.assertEqual(bloque['ramas'][nombre],
                                         [P.r1(P.num(ws.cell(r, col).value), 2)
                                          for r in range(4, len(periodos) + 4)])

    def test_seleccion_constantes_y_masa_salarial(self):
        industria = P.industria(ROOT / 'idecba/industria_ing.xlsx')
        self.assertEqual(industria['total_const'][159], 170.31)
        self.assertEqual(industria['total_const'][171], 169.62)
        self.assertEqual(industria['periodos'][159], '2015-01')
        self.assertEqual(industria['periodos'][171], '2016-01')
        for clave, fuente, fn in [('industria', FUENTES[0], P.industria),
                                   ('masa_salarial', FUENTES[1], P.masa_salarial)]:
            publicado = json.loads((ROOT / 'datos.json').read_text())[clave]
            self.assertEqual(fn(ROOT / 'idecba' / fuente), publicado)

    def test_ancla_posterior_al_inicio_y_anios_provisorios(self):
        def editar(ws):
            ws['A4'] = None
            ws['A7'] = '2002*'
        periodos, _ = self.parsear_modificado(editar)
        self.assertEqual(periodos[:4], ['2001-10', '2001-11', '2001-12', '2002-01'])

    def test_formato_ancla_en_enero(self):
        def editar(ws):
            for anio in range(2016, 2021):
                r = 180 + (anio - 2016) * 12
                ws.cell(r, 1).value = None
                ws.cell(r - 5, 1).value = anio
        periodos, _ = self.parsear_modificado(editar)
        V.validar_periodos_mensuales(periodos)
        self.assertEqual(periodos[171], '2016-01')

    def test_rechaza_contradicciones_y_formatos_rotos(self):
        casos = [({'A180': 2017}, 'contradictoria'),
                 ({'A180': 'año desconocido'}, 'año inválido'),
                 ({'B175': 'Febrero'}, 'discontinua'),
                 ({'B175': 'Diciembre'}, 'discontinua'),
                 ({'B175': 'Mes desconocido'}, 'mes inválido'),
                 ({'C175': None, 'M175': None}, 'discontinua')]
        for celdas, error in casos:
            with self.subTest(celdas=celdas), self.assertRaisesRegex(ValueError, error):
                def editar(ws):
                    for celda, valor in celdas.items():
                        ws[celda] = valor
                self.parsear_modificado(editar)

    def test_rechaza_sin_anclas_y_anio_completo_omitido(self):
        def sin_anclas(ws):
            for r in range(4, ws.max_row + 1):
                if P.anio_de(ws.cell(r, 1).value) is not None:
                    ws.cell(r, 1).value = None
        with self.assertRaisesRegex(ValueError, 'sin anclas'):
            self.parsear_modificado(sin_anclas)
        with self.assertRaisesRegex(ValueError, 'contradictoria'):
            self.parsear_modificado(lambda ws: ws.delete_rows(175, 12))

    def test_verificador_rechaza_duplicados_orden_huecos_y_formato(self):
        for periodos, error in [(['2015-01', '2015-01'], 'duplicados'),
                                (['2016-01', '2015-12'], 'orden'),
                                (['2016-01', '2016-03'], 'discontinuos'),
                                (['2016-13'], 'inválido'),
                                (['2016-1'], 'inválido')]:
            with self.subTest(periodos=periodos), self.assertRaisesRegex(ValueError, error):
                V.validar_periodos_mensuales(periodos)
        V.validar_periodos_mensuales(['2015-12', '2016-01', '2016-02'])

    def test_verificador_principal_rechaza_bloque_heredado_roto(self):
        datos = json.loads((ROOT / 'datos.json').read_text())
        datos['industria']['periodos'][171] = '2015-01'
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'datos.json'
            path.write_text(json.dumps(datos))
            with patch.object(V, 'SALIDA', str(path)), patch.object(V, 'errores', []), \
                    patch.object(V, 'avisos', []), patch('builtins.print'):
                self.assertEqual(V.main(), 1)
                self.assertTrue(any('industria: períodos mensuales duplicados' in e
                                    for e in V.errores))


if __name__ == '__main__':
    unittest.main()
